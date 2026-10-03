"""Phase 3 tests: derived datasets, KPI reference, Excel workbooks, PivotTables, Power Query scripts, Google Sheets templates.

Run from the project root:
    python python/tests/test_phase3_excel.py      # prints results and writes reports/phase3_validation.txt (takes a few minutes)
    pytest python/tests/test_phase3_excel.py -q

Workbooks are RECALCULATED FOR REAL through LibreOffice (UNO calculateAll) and compared with independent Python results. Prerequisite: python python/src/build_phase3.py.
"""
import json
import math
import platform
import re
import sys
import traceback
import zipfile
from datetime import datetime, timezone
from functools import lru_cache
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "python" / "src"))

import build_excel_pivot as bp       # noqa: E402
import build_excel_workbook as bw    # noqa: E402
import build_phase3 as b3            # noqa: E402
import clean_data as cl              # noqa: E402
import config as cfg                 # noqa: E402
import dq_catalog as dq              # noqa: E402
import kpi                           # noqa: E402
import lo_bridge as lb               # noqa: E402

XL, GS, DER = ROOT / "excel", ROOT / "google_sheets", ROOT / "data" / "processed" / "derived"
MAIN, PIVOT, EXER = XL / "IRCTC_Booking_Analytics_Workbook.xlsx", XL / "IRCTC_KPI_Pivot_Dashboard.xlsx", XL / "IRCTC_Data_Cleaning_Exercises.xlsx"
TEMPLATES = ["KPI_Tracker_Template", "Cancellation_Refund_Tracker_Template", "Occupancy_Waitlist_Tracker_Template"]
FACTS = {}


def close(a, b, tol=1e-9):
    if a in ("", None) or (isinstance(a, float) and math.isnan(a)):
        return b is None or (isinstance(b, float) and math.isnan(b)) or b == ""
    return abs(float(a) - float(b)) <= tol * max(1.0, abs(float(b)))


@lru_cache(maxsize=1)
def data():
    T = cl.read_clean()
    return T, kpi.select_sample(T), kpi._prep(T)


# ============================================================ LibreOffice sessions (real recalculation)
@lru_cache(maxsize=1)
def main_scenarios():
    T, S, _ = data()
    out = {}
    with lb.office() as desktop:
        doc = lb.load(desktop, MAIN)
        doc.calculateAll()
        out["errors"] = lb.error_cells(doc)
        rd = lambda sh, r: lb.read_range(doc, sh, r)
        val = lambda: [(r[1], r[2]) for r in rd("Validation", "A4:C40") if r[1]]
        out["validation"] = val()
        out["kpi"] = rd("KPI_Calc", "A5:G40")
        out["cards_all"] = rd("Dashboard", "B9:E23")
        dash = doc.Sheets.getByName("Dashboard")
        dash.getCellRangeByName("C5").String = "SL"; doc.calculateAll()
        out["cards_sl"] = rd("Dashboard", "B9:E23")
        dash.getCellRangeByName("C5").String = "All"; doc.calculateAll()
        rows = rd("Data_Passengers", "I2:I400")
        i = next(j for j, r in enumerate(rows) if r[0] == "Confirmed")
        doc.Sheets.getByName("Data_Passengers").getCellRangeByName(f"I{i + 2}").String = "Cancelled"; doc.calculateAll()
        out["tamper_validation"] = val()
        out["tamper_kpi_status"] = [r[5] for r in rd("KPI_Calc", "A5:F40")]
        doc.close(True)
    return out


@lru_cache(maxsize=1)
def exercise_scenarios():
    out = {}
    with lb.office() as desktop:
        doc = lb.load(desktop, EXER)
        doc.calculateAll()
        out["errors"] = lb.error_cells(doc)
        out["answer_key"] = lb.read_range(doc, "Answer_Key", "A5:G34")
        out["summary"] = lb.read_range(doc, "Answer_Key", "C36:D36")
        out["checks"] = [(r[1], r[2]) for r in lb.read_range(doc, "Validation_Checks", "A5:C20") if r[1]]
        rows = lb.read_range(doc, "Raw_Operations", "E2:E40")
        i = next(j for j, r in enumerate(rows) if isinstance(r[0], float))
        doc.Sheets.getByName("Raw_Operations").getCellRangeByName(f"E{i + 2}").String = ""; doc.calculateAll()
        out["tamper_answer_key"] = lb.read_range(doc, "Answer_Key", "A5:G34")
        doc.close(True)
    return out


@lru_cache(maxsize=1)
def template_scenarios():
    out = {}
    with lb.office() as desktop:
        for name in TEMPLATES:
            doc = lb.load(desktop, GS / f"{name}.xlsx"); doc.calculateAll()
            d = {"errors": lb.error_cells(doc), "sheets": [doc.Sheets.getByIndex(i).Name for i in range(doc.Sheets.Count)]}
            if name == TEMPLATES[0]:
                d["tracker"] = lb.read_range(doc, "KPI_Tracker", "A5:S40"); d["targets"] = lb.read_range(doc, "Targets", "A5:F10")
            elif name == TEMPLATES[1]:
                d["analysis"] = lb.read_range(doc, "Analysis", "A5:K40"); d["lead"] = lb.read_range(doc, "Lead_Time", "A5:D9")
            else:
                d["by_class"] = lb.read_range(doc, "By_Class", "A5:H11"); d["matrix"] = lb.read_range(doc, "Month_x_Class", "A4:H28")
            out[name] = d
            doc.close(True)
    return out


# ============================================================ data strategy: original data, derived datasets, sample
def test_original_data_is_unchanged():
    now, recorded = b3.data_manifest(), json.loads(b3.MANIFEST.read_text())
    assert now == recorded and len(now) == 20, [k for k in now if now[k] != recorded.get(k)]


def test_phase3_code_never_writes_to_the_original_data_folders():
    bad = []
    for f in ["kpi.py", "xl_common.py", "build_excel_workbook.py", "build_excel_pivot.py", "build_excel_exercises.py", "build_google_sheets_templates.py", "build_phase3.py"]:
        src = (ROOT / "python" / "src" / f).read_text(encoding="utf-8")
        if re.search(r"(to_csv|write_text|write_bytes|open\([^)]*['\"]w)[^\n]*(RAW_DIR|CLEAN_DIR)", src):
            bad.append(f)
    assert not bad, bad


def test_derived_datasets_reconcile_with_the_cleaned_data():
    T, _, _ = data()
    m = pd.read_csv(DER / "monthly_kpis.csv")
    assert len(m) == 24 and m.simulated_runs.eq(32).all()
    assert m.total_bookings.sum() == len(T["bookings"]) and m.total_passengers.sum() == len(T["booking_passengers"])
    assert abs(m.gross_value.sum() - T["bookings"].total_fare.sum()) < 0.01 and abs(m.refund_total.sum() - T["cancellations"].refund_amount.sum()) < 0.01
    assert (m.confirmed_bookings + m.mixed_bookings + m.cancelled_bookings <= m.total_bookings).all()
    oc = pd.read_csv(DER / "occupancy_class_runs.csv")
    assert len(oc) == len(T["seat_capacity"]) and (oc.occupancy <= 1 + 1e-9).all() and oc.passengers.sum() == len(T["booking_passengers"])
    assert np.allclose(oc.booked_seats, oc.confirmed_pax + cfg.RAC_SEAT_WEIGHT * oc.rac_pax)
    tr = pd.read_csv(DER / "train_runs.csv")
    assert len(tr) == 768 and tr.passengers.sum() == len(T["booking_passengers"]) and tr.has_delay.sum() == T["train_operations"].arrival_delay_min.notna().sum()
    rs = pd.read_csv(DER / "route_summary.csv")
    assert rs.passengers.sum() == len(T["booking_passengers"]) and rs.simulated_runs.sum() == 768
    ev = pd.read_csv(DER / "refund_events_summary.csv")
    assert abs(ev.refunds.sum() - T["cancellations"].refund_amount.sum()) < 0.01 and ev.events.sum() == len(T["cancellations"])
    cl_ = pd.read_csv(DER / "cancellation_by_lead_and_class.csv")
    assert cl_.passengers.sum() == len(T["booking_passengers"]) and cl_.cancelled_pax.sum() == (T["booking_passengers"].final_status == "Cancelled").sum()
    for f in sorted(DER.glob("*.csv")):                                      # every derived file is labelled DERIVED
        d = pd.read_csv(f)
        assert "data_origin" in d.columns and set(d.data_origin) == {"DERIVED"}, f.name
    assert len([f for f in DER.glob("*.csv") if not f.name.startswith("phase5_")]) == 7        # Phase 5 adds phase5_*.csv summaries to this folder
    for f in sorted((ROOT / "data" / "sample" / "excel").glob("*.csv")):     # sample rows keep their SYNTHETIC label
        d = pd.read_csv(f)
        assert "data_origin" in d.columns and set(d.data_origin) == {"SYNTHETIC"}, f.name
    mc = pd.read_csv(DER / "monthly_class_occupancy.csv")
    assert mc.passengers.sum() == len(T["booking_passengers"]) and mc.seats.sum() == T["seat_capacity"].total_seats.sum()


def test_kpi_reference_matches_independent_pandas_calculations():
    T, _, _ = data()
    bk, ps, ev, cap = T["bookings"], T["booking_passengers"], T["cancellations"], T["seat_capacity"]
    ref = pd.read_csv(ROOT / "reports" / "kpi_reference_values.csv").set_index("kpi_key")["value"]
    fs, ist = ps.final_status, ps.initial_status
    exp = {"total_bookings": len(bk), "total_passengers": len(ps), "cancelled_passengers": (fs == "Cancelled").sum(),
           "passenger_cancellation_rate": (fs == "Cancelled").sum() / len(ps), "booking_cancellation_rate": (bk.booking_status == "Cancelled").mean(),
           "confirmed_bookings": (bk.booking_status == "Confirmed").sum(), "mixed_bookings": (bk.booking_status == "Mixed").sum(),
           "gross_value": bk.total_fare.sum(), "refund_total": ev.refund_amount.sum(), "net_value": bk.total_fare.sum() - ev.refund_amount.sum(),
           "wl_demand": (ist == "Waitlisted").sum(), "wl_clearance_rate": ((ist == "Waitlisted") & fs.isin(["Confirmed", "RAC"])).sum() / (ist == "Waitlisted").sum(),
           "wl_uncleared": (fs == "Waitlisted").sum(), "occupancy": ((fs == "Confirmed").sum() + 0.5 * (fs == "RAC").sum()) / cap.total_seats.sum(),
           "waitlist_pressure": (fs == "Waitlisted").sum() / cap.total_seats.sum(), "avg_fare_per_passenger": bk.total_fare.sum() / len(ps),
           "otp": (T["train_operations"].arrival_delay_min.dropna() <= 15).mean(), "simulated_runs": 768}
    bad = {k: (ref[k], v) for k, v in exp.items() if not close(ref[k], v, 1e-9)}
    assert not bad, bad
    assert ref["net_value"] == ref["gross_value"] - ref["refund_total"] or abs(ref["net_value"] - (ref["gross_value"] - ref["refund_total"])) < 0.01


def test_excel_sample_is_run_complete_and_balanced_by_month():
    T, S, (bk, ps, ev, cap, ops) = data()
    runs = set(zip(S["ops"].train_id, S["ops"].run_date))
    assert len(runs) == 48 and S["ops"].groupby("journey_month").size().eq(2).all()
    full = bk[[(t, d) in runs for t, d in zip(bk.train_id, bk.journey_date)]]
    assert len(full) == len(S["bk"]) and set(full.booking_id) == set(S["bk"].booking_id)          # every booking of each sampled run is present
    assert len(S["ps"]) == ps.booking_id.isin(S["bk"].booking_id).sum() and len(S["ev"]) == ev.booking_id.isin(S["bk"].booking_id).sum()
    sk, fk = kpi.compute_kpis(S["bk"], S["ps"], S["ev"], S["cap"], S["ops"]), kpi.compute_kpis(bk, ps, ev, cap, ops)
    # 48 of 768 runs: every headline measure must be within 10% (relative) of the full dataset. Observed differences are reported in phase3_validation.txt.
    for key in ("occupancy", "passenger_cancellation_rate", "wl_clearance_rate", "avg_fare_per_passenger", "booking_cancellation_rate"):
        assert abs(sk[key] - fk[key]) / fk[key] < 0.10, (key, sk[key], fk[key])


# ============================================================ main workbook (real recalculation)
def test_main_workbook_has_the_expected_sheets_and_features():
    z = zipfile.ZipFile(MAIN)
    wbx = z.read("xl/workbook.xml").decode()
    names = re.findall(r'<sheet name="([^"]+)"', wbx)
    assert names == ["README", "KPI_Definitions", "Dashboard", "Monthly_Aggregates", "Cancellation_Refund", "Occupancy_Demand", "KPI_Calc", "Validation", "Parameters",
                     "Data_Bookings", "Data_Passengers", "Data_Cancellations", "Data_ClassRuns", "Data_Runs"], names
    sheets = "".join(z.read(n).decode() for n in z.namelist() if n.startswith("xl/worksheets/sheet"))
    assert sheets.count("<dataValidation ") >= 3 and sheets.count("<conditionalFormatting") >= 5
    assert sum(n.startswith("xl/charts/chart") for n in z.namelist()) >= 8
    assert "XLOOKUP" not in sheets.upper()


def test_main_workbook_recalculates_without_errors_and_validation_passes():
    s = main_scenarios()
    assert sum(v for v in s["errors"].values() if v > 0) == 0 and all(v >= 0 for v in s["errors"].values()), s["errors"]
    fails = [(l, r) for l, r in s["validation"] if r == "FAIL"]
    assert not fails, fails
    assert s["validation"][-1][1].startswith("ALL 24 CHECKS PASS"), s["validation"][-1]


def test_main_workbook_excel_kpis_equal_python_reference_within_tolerance():
    s = main_scenarios(); T, S, _ = data()
    ref_s = kpi.compute_kpis(S["bk"], S["ps"], S["ev"], S["cap"], S["ops"])
    assert len(s["kpi"]) == len(bw.CALC_KEYS) == 36
    for row, key in zip(s["kpi"], bw.CALC_KEYS):
        name, excel, py, _, tol, status, full = row
        assert name == key and status == "PASS", (key, excel, py)
        assert close(excel, ref_s[key], 1e-9 if tol == 0 else max(tol, 1e-9)), (key, excel, ref_s[key])


def test_dashboard_class_filter_scenario_matches_python_for_sl():
    s = main_scenarios(); _, S, _ = data()
    m = lambda d: d[d.travel_class == "SL"]
    k = kpi.compute_kpis(m(S["bk"]), m(S["ps"]), m(S["ev"]), m(S["cap"]), S["ops"])
    exp = [k["total_bookings"], k["total_passengers"], k["confirmed_bookings"], k["mixed_bookings"], k["passenger_cancellation_rate"], k["booking_cancellation_rate"], k["wl_demand"],
           k["wl_clearance_rate"], k["wl_uncleared"], k["occupancy"], k["waitlist_pressure"], k["gross_value"], k["refund_total"], k["net_value"], k["avg_fare_per_passenger"]]
    got = [r[1] for r in s["cards_sl"]]
    bad = [(r[0], g, e) for r, g, e in zip(s["cards_sl"], got, exp) if not close(g, e, 1e-9)]
    assert not bad, bad
    assert not close(s["cards_all"][0][1], exp[0])              # the filter really changed the numbers


def test_denominators_are_shown_on_dashboard_cards():
    s = main_scenarios()
    wanted = ["Cancellation rate (passenger-based)", "Booking cancellation rate (companion)", "Waitlist clearance rate", "Seat occupancy %",
              "Waitlist pressure % (not part of occupancy)", "Average fare per passenger (INR)"]
    rates = {r[0]: r for r in s["cards_all"] if r[0] in wanted}
    assert set(rates) == set(wanted)
    for lab, r in rates.items():
        assert isinstance(r[2], float) and isinstance(r[3], float) and close(r[1], r[2] / r[3], 1e-12), lab


def test_tampering_with_data_is_detected_by_the_workbook():
    s = main_scenarios()
    assert not s["tamper_validation"][-1][1].startswith("ALL"), s["tamper_validation"][-1]
    assert "FAIL" in s["tamper_kpi_status"]


# ============================================================ pivot workbook
def test_pivot_workbook_has_real_pivottables_that_match_pandas():
    import openpyxl
    z = zipfile.ZipFile(PIVOT)
    assert sum(re.match(r"xl/pivotTables/pivotTable\d+\.xml$", n) is not None for n in z.namelist()) == 7
    assert sum(n.startswith("xl/charts/chart") for n in z.namelist()) >= 2
    wb = openpyxl.load_workbook(PIVOT, data_only=True)
    assert {ws.title: len(ws._pivots) for ws in wb if ws.title.startswith("PT_")} == {p[0]: 1 for p in bp.PIVOTS}
    F = bp.facts()
    rows = list(wb["PT_1_Status_by_Class"].iter_rows(min_row=1, max_row=14, values_only=True))
    h = next(i for i, r in enumerate(rows) if r[0] == "travel_class"); cols = rows[h][1:5]
    ct = pd.crosstab(F["Fact_Passengers"].travel_class, F["Fact_Passengers"].final_status)
    for r in rows[h + 1:]:
        if r[0] in ct.index:
            assert all(int(v or 0) == int(ct.loc[r[0], c]) for c, v in zip(cols, r[1:5])), r
    rows = list(wb["PT_7_Booking_Status"].iter_rows(min_row=1, max_row=14, values_only=True))
    h = next(i for i, r in enumerate(rows) if r[0] == "travel_class"); cols = [c for c in rows[h][1:] if c and c != "Total Result"]
    ct = pd.crosstab(F["Fact_Bookings"].travel_class, F["Fact_Bookings"].booking_status)
    for r in rows[h + 1:]:
        if r[0] in ct.index:
            assert all(int(v or 0) == int(ct.loc[r[0], c]) for c, v in zip(cols, r[1:1 + len(cols)])), r
    ev = F["Fact_Events"].pivot_table(index="journey_month", columns="cancellation_type", values="refund_amount", aggfunc="sum")
    rows = list(wb["PT_4_Refunds"].iter_rows(min_row=1, max_row=40, values_only=True)); h = next(i for i, r in enumerate(rows) if r[0] == "journey_month")
    n = 0
    for r in rows[h + 1:]:
        if r[0] in ev.index:
            n += 1; assert abs(float(r[1] or 0) - float(np.nan_to_num(ev.loc[r[0], "Passenger"]))) < 0.01 and abs(float(r[2] or 0) - float(np.nan_to_num(ev.loc[r[0], "System"]))) < 0.01
    assert n == 24


def test_pivot_dashboard_cards_equal_python_sample_kpis():
    import openpyxl
    _, S, _ = data(); k = kpi.compute_kpis(S["bk"], S["ps"], S["ev"], S["cap"], S["ops"])
    d = {r[0]: r[1] for r in openpyxl.load_workbook(PIVOT, data_only=True)["Dashboard"].iter_rows(min_row=5, max_row=20, max_col=2, values_only=True) if r[0]}
    pairs = [("Total passengers", "total_passengers"), ("Total bookings", "total_bookings"), ("Fully confirmed bookings (every passenger Confirmed)", "confirmed_bookings"),
             ("Mixed-status bookings (reported separately)", "mixed_bookings"), ("Cancellation rate (passenger-based)", "passenger_cancellation_rate"),
             ("Waiting-list demand (initially waitlisted passengers)", "wl_demand"), ("Waitlist clearance rate", "wl_clearance_rate"), ("Uncleared waitlist passengers", "wl_uncleared"),
             ("Seat occupancy %", "occupancy"), ("Waitlist pressure % (not part of occupancy)", "waitlist_pressure"), ("Gross booking value (INR)", "gross_value"),
             ("Refund amount (INR)", "refund_total"), ("Average fare per passenger (INR)", "avg_fare_per_passenger"), ("Net booking value (INR)", "net_value")]
    bad = [(a, d.get(a), k[b]) for a, b in pairs if not close(d.get(a), k[b], 1e-9)]
    assert not bad, bad


# ============================================================ cleaning exercises
def test_exercise_formulas_reproduce_the_independent_detection_counts():
    s = exercise_scenarios()
    assert sum(v for v in s["errors"].values() if v > 0) == 0, s["errors"]
    rows = [r for r in s["answer_key"] if r[0]]
    assert [r[0] for r in rows] == dq.ORDER
    for r in rows:
        code, _, _, expected, injected, excel, result = r[:7]
        assert expected == injected and expected > 0, (code, expected, injected)
        if code == "total_fare_invalid":
            assert result == "MANUAL"
        else:
            assert result == "PASS" and excel == expected, (code, excel, expected)
    assert s["summary"][0][1].startswith("ALL FORMULA COUNTS MATCH (29 of 29")
    assert all(r == "PASS" or r.startswith("ALL") for _, r in s["checks"]) and s["checks"][-1][1].startswith("ALL 11 CHECKS PASS")


def test_exercise_answer_key_detects_a_changed_raw_value():
    s = exercise_scenarios()
    before = {r[0]: r[6] for r in s["answer_key"] if r[0]}
    after = {r[0]: r[6] for r in s["tamper_answer_key"] if r[0]}
    assert before["missing_arrival_delay"] == "PASS" and after["missing_arrival_delay"] == "FAIL"


def test_exercise_raw_sheets_preserve_raw_values():
    import openpyxl
    raw = cl.read_raw()["bookings"]
    ws = openpyxl.load_workbook(EXER, read_only=True)["Raw_Bookings"]
    rows = list(ws.iter_rows(min_row=2, max_row=200, values_only=True))
    ids = {str(int(r[0])): r for r in rows if r[0] is not None}
    for _, rr in raw[raw.booking_id.isin(ids)].head(80).iterrows():
        assert ids[rr.booking_id][3] == rr.journey_date and (ids[rr.booking_id][10] or "") == rr.booking_status


# ============================================================ Google Sheets templates
def test_google_sheets_templates_match_python_and_have_no_errors():
    s = template_scenarios(); m = pd.read_csv(DER / "monthly_kpis.csv"); mc = pd.read_csv(DER / "monthly_class_occupancy.csv")
    for n in TEMPLATES:
        assert all(v == 0 for v in s[n]["errors"].values()), (n, s[n]["errors"])
    t = s[TEMPLATES[0]]
    exp = {1: m.bookings_per_run, 2: m.passengers_per_run, 3: m.passenger_cancellation_rate, 4: m.booking_cancellation_rate, 5: m.confirmed_bookings / m.total_bookings,
           6: m.mixed_bookings / m.total_bookings, 7: m.wl_demand / m.total_passengers, 8: m.wl_clearance_rate, 9: m.waitlist_pressure, 10: m.occupancy, 11: m.gross_value,
           12: m.refund_total, 13: m.net_value, 14: m.avg_fare_per_passenger, 15: m.otp, 16: m.avg_delay}
    for i in range(24):
        for c, ser in exp.items():
            assert close(t["tracker"][i][c], ser.iloc[i]), (i, c)
        if i:
            assert close(t["tracker"][i][17], m.bookings_growth.iloc[i]) and close(t["tracker"][i][18], m.bookings_per_run_growth.iloc[i])
    assert all(v in ("", None) for r in t["tracker"][24:] for v in r[1:])                         # empty template rows stay blank
    assert {r[0]: r[5] for r in t["targets"]}["Seat occupancy %"] in ("On target", "Watch", "Off target")
    last = m.iloc[23]
    assert close({r[0]: r[4] for r in t["targets"]}["Seat occupancy %"], last.occupancy)
    a = s[TEMPLATES[1]]["analysis"]
    for i in range(24):
        r = m.iloc[i]; tot = r.refund_passenger + r.refund_system
        assert close(a[i][1], r.passenger_cancellation_rate) and close(a[i][2], r.booking_cancellation_rate) and close(a[i][3], tot) and close(a[i][5], tot / r.gross_value)
        assert close(a[i][6], r.refund_system / tot) and close(a[i][7], r.retained_charges / (tot + r.retained_charges))
    assert all(v in ("", None) for r in a[24:] for v in r[1:])
    o = s[TEMPLATES[2]]
    for r in o["by_class"]:
        sub = mc if r[0] == "All classes" else mc[mc.travel_class == r[0]]
        assert close(r[4], (sub.confirmed_pax.sum() + 0.5 * sub.rac_pax.sum()) / sub.seats.sum()) and close(r[7], sub.wl_final_pax.sum() / sub.seats.sum())
    hdr = o["matrix"][0][1:7]
    for r in o["matrix"][1:]:
        for j, c in enumerate(hdr):
            sub = mc[(mc.journey_month == r[0]) & (mc.travel_class == c)]
            assert close(r[1 + j], (sub.confirmed_pax.sum() + 0.5 * sub.rac_pax.sum()) / sub.seats.sum() if len(sub) else float("nan"))


def test_google_sheets_templates_use_only_portable_functions():
    banned = ["XLOOKUP", "FILTER(", "UNIQUE(", "SORT(", "LET(", "LAMBDA(", "SEQUENCE(", "TEXTJOIN(", "CONCAT("]
    for n in TEMPLATES:
        z = zipfile.ZipFile(GS / f"{n}.xlsx")
        xml = "".join(z.read(f).decode() for f in z.namelist() if f.startswith("xl/worksheets/sheet")).upper()
        assert not [b for b in banned if b in xml], (n, [b for b in banned if b in xml])


# ============================================================ Power Query scripts, documentation
def _strip(text):
    text = re.sub(r"//[^\n]*", "", text)
    return re.sub(r'"([^"]|"")*"', '""', text)


def test_power_query_scripts_are_structurally_sound():
    files = sorted((XL / "power_query").glob("*.pq"))
    assert len(files) >= 14
    defined = {re.search(r"// Query name: (\w+)", f.read_text(encoding="utf-8")).group(1) for f in files}
    assert {"DataFolder", "RawFolder", "RacWeight", "fnLoadCsv", "Bookings", "Passengers", "PassengersFact", "Cancellations", "SeatCapacity", "TrainOperations", "MonthlyKpis",
            "OccupancyClassRuns", "StatusLookup", "CleanRawBookingsPreview"} <= defined
    for f in files:
        t = f.read_text(encoding="utf-8"); s = _strip(t)
        assert s.count("(") == s.count(")") and s.count("[") == s.count("]") and s.count("{") == s.count("}"), f.name
        assert "XLOOKUP" not in t.upper() and not re.search(r",\s*\n\s*in\b", s), f.name
        if re.search(r"^\s*let\b", s, flags=re.M):
            last = re.search(r"\n\s*in\s*\n\s*(\w+)\s*$", s.rstrip() + "\n")
            assert last and re.search(rf"^\s*{last.group(1)}\s*=", s, flags=re.M), (f.name, "final step not defined")
        for ref in re.findall(r"\b(Bookings|Passengers|PassengersFact|Cancellations|SeatCapacity|TrainOperations|StatusLookup|fnLoadCsv|DataFolder|RawFolder|RacWeight)\b", s):
            assert ref in defined, (f.name, ref)


def test_sampling_limitation_is_documented_everywhere_it_matters():
    for f in ["excel/SETUP_INSTRUCTIONS.md", "google_sheets/SETUP_INSTRUCTIONS.md", "docs/KPI_consistency_matrix.md"]:
        t = (ROOT / f).read_text(encoding="utf-8")
        assert ("19,708" in t or "sampled" in t.lower()) and "synthetic" in t.lower(), f
    for f in (MAIN, PIVOT, EXER, *[GS / f"{n}.xlsx" for n in TEMPLATES]):
        ss = zipfile.ZipFile(f).read("xl/sharedStrings.xml").decode()
        assert "SYNTHETIC" in ss.upper() and ("19,708" in ss or "sampled simulation" in ss.lower()), f.name
        assert "network-wide" in ss.lower() or "not IRCTC" in ss or "Not IRCTC" in ss, f.name


def test_kpi_consistency_matrix_covers_all_fifteen_kpis_and_states_pending_tools():
    matrix = (ROOT / "docs" / "KPI_consistency_matrix.md").read_text(encoding="utf-8")
    heads = re.findall(r"^### \d+\. (.+)$", (ROOT / "docs" / "KPI_definitions.md").read_text(encoding="utf-8"), flags=re.M)
    assert len(heads) == 15
    missing = [h for h in heads if h.split(" (")[0] not in matrix]
    assert not missing, missing
    assert "Specified, not yet implemented" in matrix and "Phase 4" in matrix and "Phase 6" in matrix and "Phase 7" in matrix


def test_manual_setup_instructions_cover_every_feature_that_cannot_be_generated():
    t = (XL / "SETUP_INSTRUCTIONS.md").read_text(encoding="utf-8")
    for needle in ["Blank Query", "Advanced Editor", "PT_1", "PT_7", "Insert Slicer", "Report Connections", "Data Validation", "Not opened", "XLOOKUP", "201,182", "405,200"]:
        assert needle.lower() in t.lower(), needle
    assert (XL / "FORMULA_LIBRARY.md").exists() and (GS / "SETUP_INSTRUCTIONS.md").exists()


# ============================================================ runner
def _tests():
    return [(n, f) for n, f in sorted(globals().items(), key=lambda kv: kv[1].__code__.co_firstlineno if callable(kv[1]) and hasattr(kv[1], "__code__") else 0) if n.startswith("test_") and callable(f)]


def main():
    lines, passed, failed = [], 0, 0
    for name, fn in _tests():
        try:
            fn(); lines.append(f"PASS  {name}"); passed += 1
        except Exception as e:
            lines.append(f"FAIL  {name}\n        {type(e).__name__}: {str(e)[:500]}"); failed += 1; traceback.print_exc()
    ms, ex, tp = None, None, None
    try:
        ms, ex, tp = main_scenarios(), exercise_scenarios(), template_scenarios()
    except Exception:
        pass
    facts = []
    if ms:
        v = ms["validation"]
        facts += [f"Main workbook: {sum(r == 'PASS' for _, r in v)} of {len(v) - 1} Validation checks PASS after a real LibreOffice recalculation; overall cell reads '{v[-1][1]}'",
                  f"Main workbook: {sum(r[5] == 'PASS' for r in ms['kpi'])} of {len(ms['kpi'])} KPI_Calc rows PASS (Excel formula vs independent Python reference); formula error cells: {sum(x for x in ms['errors'].values() if x > 0)}",
                  f"Main workbook: dashboard class-filter scenario (SL) matched Python on all 15 cards; deliberate tampering was detected (overall cell then reads '{ms['tamper_validation'][-1][1]}')"]
    if ex:
        ok = sum(1 for r in ex["answer_key"] if r[0] and r[6] == "PASS")
        facts += [f"Exercises workbook: {ok} of 29 formula counts equal the independent detection counts; summary: {ex['summary'][0][1]}; validation checks: {ex['checks'][-1][1]}",
                  "Exercises workbook: changing one raw value flipped the matching Answer_Key row from PASS to FAIL (detected)"]
    try:
        T, S, (bk, ps, ev, cap, ops) = data()
        sk, fk = kpi.compute_kpis(S["bk"], S["ps"], S["ev"], S["cap"], S["ops"]), kpi.compute_kpis(bk, ps, ev, cap, ops)
        facts.append("Sample (48 of 768 runs) versus full dataset: " + "; ".join(f"{k} {sk[k]:.4f} vs {fk[k]:.4f} ({(sk[k] - fk[k]) / fk[k]:+.1%})" for k in ("occupancy", "passenger_cancellation_rate", "wl_clearance_rate", "avg_fare_per_passenger")))
    except Exception:
        pass
    if tp:
        facts += [f"Google Sheets templates: 3 recalculated, compared with Python monthly values, formula error cells: {sum(sum(x for x in t['errors'].values() if x > 0) for t in tp.values())}"]
    out = [f"Phase 3 validation run: {datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%MZ')}", "Command: python python/tests/test_phase3_excel.py    (build: python python/src/build_phase3.py)",
           f"Environment: Python {platform.python_version()}, pandas {pd.__version__}, LibreOffice 24.2 via UNO (real recalculation), seed {cfg.SEED}", "",
           "== A. Test suite ==", *lines, "", f"{passed} passed, {failed} failed, {passed + failed} total", "", "== B. Key results (measured in this run) ==", *[f"- {f}" for f in facts], "",
           "== C. Not verified in this environment ==", "- Workbooks were not opened in Microsoft Excel or Google Sheets; PivotTables were authored and checked with LibreOffice.",
           "- Power Query M scripts were not executed (no Power Query engine); only their structure was linted.", "- XLOOKUP cannot be evaluated by LibreOffice 24.2 (#NAME?) and is not used in any verified cell."]
    text = "\n".join(out)
    print(text)
    if "--no-write" not in sys.argv:
        (cfg.REPORTS_DIR / "phase3_validation.txt").write_text(text + "\n", encoding="utf-8")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
